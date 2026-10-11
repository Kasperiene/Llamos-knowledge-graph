"""
Batch runner — process multiple PDFs and build a combined knowledge graph.

Usage:
    python batch_runner.py
    python batch_runner.py data/papers/odetos_straipsniai/
    python batch_runner.py data/papers/odetos_straipsniai/ --excel
    python batch_runner.py data/papers/odetos_straipsniai/ --limit 3
    python batch_runner.py data/papers/odetos_straipsniai/ --no-cache --excel
"""

import argparse
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

# ConceptRegistry and normalisation now live in Agent 5
from agents.agent5_entity_linking import ConceptRegistry, normalize as _normalize

# ─────────────────────────────────────────── graph builder ───────────────────

def build_graph(papers_data: list, registry: ConceptRegistry,
                all_relations: list):
    import networkx as nx
    G = nx.DiGraph()

    # Concept nodes
    for c in registry.concepts():
        G.add_node(c["norm_key"],
                   node_type="Concept",
                   label=c["display_text"],
                   concept_type=c["type"],
                   frequency=c["frequency"],
                   paper_count=c["paper_count"])

    # Paper + Author nodes
    for pd in papers_data:
        pid = pd["paper_id"]
        G.add_node(pid,
                   node_type="Paper",
                   label=(pd["title"] or pd["filename"])[:80],
                   title=pd["title"],
                   doi=pd.get("doi", ""),
                   word_count=pd.get("word_count", 0))

        for norm_key in pd.get("concept_keys", []):
            if G.has_node(norm_key):
                G.add_edge(pid, norm_key, rel="mentions")

        for author in pd.get("authors", []):
            aid = f"author::{_normalize(author)}"
            if not G.has_node(aid):
                G.add_node(aid, node_type="Author", label=author)
            G.add_edge(aid, pid, rel="authored")

    # Concept–Concept relation edges
    canon = registry.canon_keys()
    for r in all_relations:
        h = _normalize(r["head"])
        t = _normalize(r["tail"])
        if h in canon and t in canon:
            if G.has_edge(h, t):
                G[h][t]["weight"] = G[h][t].get("weight", 1) + 1
            else:
                G.add_edge(h, t,
                           rel=r.get("type", "related"),
                           source=r.get("source", "svo"),
                           weight=1)

    return G


# ─────────────────────────────────────────── Excel summary ───────────────────

def write_excel(papers_data: list, registry: ConceptRegistry,
                graph, out_path: Path,
                all_relations: list | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side)

    wb = Workbook()

    GOLD_BG  = "FFF3B0"
    BLUE_BG  = "DCE6F1"
    BLUE_H   = "1F497D"
    ALT_ROW  = "F9F9F9"

    hdr_fill  = PatternFill("solid", fgColor=BLUE_BG)
    gold_fill = PatternFill("solid", fgColor=GOLD_BG)
    alt_fill  = PatternFill("solid", fgColor=ALT_ROW)
    hdr_font  = Font(bold=True, color=BLUE_H)
    title_font= Font(bold=True, size=13, color=BLUE_H)
    thin      = Side(border_style="thin", color="CCCCCC")
    brd       = Border(left=thin, right=thin, top=thin, bottom=thin)

    def H(ws, row, cols, labels):
        for col, label in zip(cols, labels):
            c = ws.cell(row=row, column=col, value=label)
            c.font = hdr_font; c.fill = hdr_fill; c.border = brd
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    def C(ws, row, col, val, fill=None, bold=False, wrap=False):
        c = ws.cell(row=row, column=col, value=val)
        c.border = brd
        c.alignment = Alignment(horizontal="left", wrap_text=wrap)
        if fill: c.fill = fill
        if bold: c.font = Font(bold=True)
        return c

    def alt(i):
        return alt_fill if i % 2 else None

    # ── Sheet 1: Papers Overview ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Papers Overview"
    ws1.cell(1, 1, "Knowledge Graph — Batch Processing Summary").font = title_font
    H(ws1, 2, range(1, 9),
      ["#", "Filename", "Title", "Authors", "DOI", "Words", "Concepts", "Relations"])
    for w, col in zip([5, 42, 45, 30, 22, 10, 12, 12], "ABCDEFGH"):
        ws1.column_dimensions[col].width = w

    for i, p in enumerate(papers_data):
        r = i + 3
        C(ws1, r, 1, i + 1,         fill=alt(i))
        C(ws1, r, 2, p["filename"],  fill=alt(i))
        C(ws1, r, 3, p["title"] or "—", fill=alt(i), wrap=True)
        C(ws1, r, 4, ", ".join(p.get("authors", [])[:3]), fill=alt(i))
        C(ws1, r, 5, p.get("doi", ""),   fill=alt(i))
        C(ws1, r, 6, p.get("word_count", 0), fill=alt(i))
        C(ws1, r, 7, len(p.get("concept_keys", [])), fill=alt(i))
        C(ws1, r, 8, p.get("relation_count", 0),     fill=alt(i))

    # ── Sheet 2: Top Concepts ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Top Concepts")
    ws2.cell(1, 1, "Concepts ranked by paper coverage and frequency").font = title_font
    H(ws2, 2, range(1, 7),
      ["Rank", "Concept", "Type", "Papers (N)", "Total Freq", "Best Score"])
    for w, col in zip([8, 42, 14, 12, 12, 12], "ABCDEF"):
        ws2.column_dimensions[col].width = w

    sorted_concepts = sorted(
        registry.concepts(),
        key=lambda c: (c["paper_count"], c["frequency"]),
        reverse=True,
    )
    for i, c in enumerate(sorted_concepts[:150]):
        r = i + 3
        fill = gold_fill if c["paper_count"] >= 5 else alt(i)
        C(ws2, r, 1, i + 1,                fill=fill)
        C(ws2, r, 2, c["display_text"],     fill=fill)
        C(ws2, r, 3, c["type"],             fill=fill)
        C(ws2, r, 4, c["paper_count"],      fill=fill, bold=c["paper_count"] >= 5)
        C(ws2, r, 5, c["frequency"],        fill=fill)
        C(ws2, r, 6, round(c["best_score"], 3), fill=fill)

    # ── Sheet 3: Co-occurrence ────────────────────────────────────────────────
    ws3 = wb.create_sheet("Co-occurrence")
    ws3.cell(1, 1, "Concept pairs co-occurring across papers (≥ 2 shared)").font = title_font
    H(ws3, 2, range(1, 5), ["Concept A", "Concept B", "Shared Papers", "Paper IDs (short)"])
    for w, col in zip([38, 38, 15, 55], "ABCD"):
        ws3.column_dimensions[col].width = w

    multi = [c for c in registry.concepts() if c["paper_count"] >= 3]
    cooc = []
    for i, ca in enumerate(multi):
        for cb in multi[i + 1:]:
            shared = set(ca["papers"]) & set(cb["papers"])
            if len(shared) >= 2:
                cooc.append((ca["display_text"], cb["display_text"],
                             len(shared), sorted(shared)))
    cooc.sort(key=lambda x: x[2], reverse=True)

    for i, (a, b, cnt, pids) in enumerate(cooc[:100]):
        r = i + 3
        fill = gold_fill if cnt >= 5 else alt(i)
        C(ws3, r, 1, a,   fill=fill)
        C(ws3, r, 2, b,   fill=fill)
        C(ws3, r, 3, cnt, fill=fill, bold=cnt >= 5)
        short_ids = [p[:30] for p in pids[:6]]
        C(ws3, r, 4, ", ".join(short_ids), fill=fill, wrap=True)

    # ── Sheet 4: Per-Paper Concepts ───────────────────────────────────────────
    ws4 = wb.create_sheet("Per-Paper Concepts")
    ws4.cell(1, 1, "Top concepts extracted per paper").font = title_font
    H(ws4, 2, range(1, 5), ["Paper (short)", "Concept", "Type", "Score"])
    for w, col in zip([42, 38, 14, 12], "ABCD"):
        ws4.column_dimensions[col].width = w

    r = 3
    for p in papers_data:
        short = p["filename"][:40]
        for c in p.get("top_concepts", [])[:25]:
            fill = alt(r)
            C(ws4, r, 1, short,         fill=fill)
            C(ws4, r, 2, c["text"],     fill=fill)
            C(ws4, r, 3, c["type"],     fill=fill)
            C(ws4, r, 4, round(c["score"], 3), fill=fill)
            r += 1

    # ── Sheet 5: Graph Stats ──────────────────────────────────────────────────
    ws5 = wb.create_sheet("Graph Stats")
    ws5.cell(1, 1, "Knowledge Graph Statistics").font = title_font
    H(ws5, 2, [1, 2], ["Metric", "Value"])
    ws5.column_dimensions["A"].width = 38
    ws5.column_dimensions["B"].width = 20

    n_concepts   = sum(1 for _, d in graph.nodes(data=True) if d.get("node_type") == "Concept")
    n_papers_g   = sum(1 for _, d in graph.nodes(data=True) if d.get("node_type") == "Paper")
    n_authors_g  = sum(1 for _, d in graph.nodes(data=True) if d.get("node_type") == "Author")
    n_rel_edges  = sum(1 for _, _, d in graph.edges(data=True)
                       if d.get("rel") not in ("mentions", "authored"))
    cross_paper  = sum(1 for c in registry.concepts() if c["paper_count"] > 1)

    stats = [
        ("PDFs processed",              len(papers_data)),
        ("Total unique concepts",        len(registry.concepts())),
        ("Cross-paper concepts",         cross_paper),
        ("Graph: total nodes",           graph.number_of_nodes()),
        ("Graph: total edges",           graph.number_of_edges()),
        ("Graph: Paper nodes",           n_papers_g),
        ("Graph: Concept nodes",         n_concepts),
        ("Graph: Author nodes",          n_authors_g),
        ("Graph: Concept–Concept edges", n_rel_edges),
    ]
    for i, (metric, val) in enumerate(stats):
        C(ws5, i + 3, 1, metric, fill=alt(i))
        C(ws5, i + 3, 2, val,    fill=alt(i), bold=True)

    # ── Sheet 6: Relation Triplets ────────────────────────────────────────────
    if all_relations:
        ws6 = wb.create_sheet("Relation Triplets")
        ws6.cell(1, 1, "Subject → Relation → Object triplets extracted by Agent 4 (SVO)").font = title_font
        H(ws6, 2, range(1, 7),
          ["#", "Head (Subject)", "Relation (Verb)", "Tail (Object)", "Source", "Paper"])
        for w, col in zip([7, 38, 22, 38, 10, 42], "ABCDEF"):
            ws6.column_dimensions[col].width = w

        # build a paper-id → short filename lookup
        pid_to_name = {p["paper_id"]: p["filename"][:40] for p in papers_data}

        # deduplicate by (head_lower, rel, tail_lower) and count occurrences
        from collections import Counter as _Counter
        triple_counts: dict[tuple, dict] = {}
        for p in papers_data:
            for r in p.get("relations", []):
                key = (r["head"].lower().strip(), r["type"].strip(), r["tail"].lower().strip())
                if key not in triple_counts:
                    triple_counts[key] = {
                        "head":    r["head"].strip(),
                        "type":    r["type"].strip(),
                        "tail":    r["tail"].strip(),
                        "source":  r.get("source", "svo"),
                        "papers":  set(),
                        "count":   0,
                    }
                triple_counts[key]["papers"].add(p["paper_id"])
                triple_counts[key]["count"] += 1

        # sort: cross-paper first, then by count
        sorted_triples = sorted(
            triple_counts.values(),
            key=lambda t: (len(t["papers"]), t["count"]),
            reverse=True,
        )

        for i, t in enumerate(sorted_triples):
            r_idx = i + 3
            is_cross = len(t["papers"]) > 1
            fill = gold_fill if is_cross else alt(i)
            paper_names = ", ".join(pid_to_name.get(pid, pid[:30]) for pid in sorted(t["papers"])[:3])
            C(ws6, r_idx, 1, i + 1,          fill=fill)
            C(ws6, r_idx, 2, t["head"],       fill=fill)
            C(ws6, r_idx, 3, t["type"],       fill=fill)
            C(ws6, r_idx, 4, t["tail"],       fill=fill)
            C(ws6, r_idx, 5, t["source"],     fill=fill)
            C(ws6, r_idx, 6, paper_names,     fill=fill, wrap=True)

    wb.save(str(out_path))
    print(f"  Excel → {out_path}")


# ─────────────────────────────────────────── per-paper processing ────────────

def process_paper(pdf_path: Path, out_dir: Path,
                  agent1, agent3, agent4,
                  no_cache: bool = False) -> dict | None:
    safe_stem = re.sub(r"[^a-zA-Z0-9_-]", "_", pdf_path.stem[:60])
    cache_file = out_dir / f"{safe_stem}.json"

    if not no_cache and cache_file.exists():
        print(f"  [cache] {pdf_path.name}")
        return json.loads(cache_file.read_text())

    print(f"\n{'─'*60}")
    print(f"  File: {pdf_path.name}")
    t0 = time.perf_counter()

    try:
        ingestion = agent1.run(str(pdf_path))
    except Exception as e:
        print(f"  [SKIP] Ingestion failed: {e}")
        print(f"         (PDF may be scanned/image-only — install poppler for OCR support)")
        return None

    if not ingestion.get("text", "").strip():
        print(f"  [SKIP] No text extracted from {pdf_path.name}")
        return None

    try:
        concepts  = agent3.run(ingestion["text"])
    except Exception as e:
        print(f"  [SKIP] Concept extraction failed: {e}")
        return None

    try:
        relations = agent4.run(ingestion["text"], concepts)
    except Exception as e:
        print(f"  [WARN] Relation extraction failed: {e}")
        relations = []

    elapsed = time.perf_counter() - t0

    result = {
        "paper_id":       safe_stem,
        "filename":       pdf_path.name,
        "title":          ingestion.get("title", ""),
        "authors":        ingestion.get("authors", []),
        "doi":            ingestion.get("doi", ""),
        "abstract":       ingestion.get("abstract", ""),
        "word_count":     len(ingestion["text"].split()),
        "sections":       ingestion.get("sections", []),
        "top_concepts":   concepts[:50],
        "relations":      relations,
        "concept_count":  len(concepts),
        "relation_count": len(relations),
        "elapsed_s":      round(elapsed, 1),
    }

    cache_file.write_text(json.dumps(result, indent=2, ensure_ascii=False))
    print(f"  Done in {elapsed:.1f}s  concepts={len(concepts)}  relations={len(relations)}")
    return result


# ─────────────────────────────────────────── main ────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="Batch-process PDFs into a knowledge graph")
    ap.add_argument("folder",    nargs="?",
                    default="data/papers/odetos_straipsniai",
                    help="Folder containing PDF files")
    ap.add_argument("--out-dir", default="data/output/batch",
                    help="Directory for per-paper JSON cache files")
    ap.add_argument("--graph-out", default="data/output/knowledge_graph",
                    help="Base path for graph output (no extension)")
    ap.add_argument("--excel",    action="store_true",
                    help="Generate batch_summary.xlsx after processing")
    ap.add_argument("--no-cache", action="store_true",
                    help="Ignore existing cached JSON and re-process")
    ap.add_argument("--limit",   type=int, default=0,
                    help="Process only the first N PDFs (0 = all)")
    args = ap.parse_args()

    paper_dir = Path(args.folder)
    if not paper_dir.exists():
        print(f"Error: folder not found: {paper_dir}")
        sys.exit(1)

    pdf_files = sorted(paper_dir.glob("*.pdf"))
    if not pdf_files:
        print(f"No PDFs found in {paper_dir}")
        sys.exit(1)

    if args.limit:
        pdf_files = pdf_files[:args.limit]

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n{'='*60}")
    print(f"  Batch runner")
    print(f"  Folder  : {paper_dir}")
    print(f"  PDFs    : {len(pdf_files)}")
    print(f"  Out dir : {out_dir}")
    print(f"{'='*60}")

    # Initialise agents once — models are loaded lazily on first call
    sys.path.insert(0, str(Path(__file__).parent))
    from agents.agent1_ingestion import DataIngestionAgent
    from agents.agent3_concept_extraction import ConceptExtractionAgent
    from agents.agent4_relation_extraction import RelationExtractionAgent

    agent1 = DataIngestionAgent()
    agent3 = ConceptExtractionAgent(
        spacy_model="en_core_sci_sm",          # SciSpacy — best for scientific NER
        embedding_model="BAAI/bge-small-en-v1.5",  # top KeyBERT quality from benchmarks
        top_n_keybert=30,
        top_n_yake=30,
    )
    agent4 = RelationExtractionAgent(spacy_model="en_core_web_sm")

    registry   = ConceptRegistry()
    papers_data = []
    t_total    = time.perf_counter()

    skipped = []
    for pdf_path in pdf_files:
        result = process_paper(pdf_path, out_dir, agent1, agent3, agent4,
                               no_cache=args.no_cache)
        if result is None:
            skipped.append(pdf_path.name)
            continue
        concept_keys = [
            registry.add(c, result["paper_id"])
            for c in result.get("top_concepts", [])
        ]
        result["concept_keys"] = concept_keys
        papers_data.append(result)

    if skipped:
        print(f"\n  Skipped {len(skipped)} PDF(s): {', '.join(skipped)}")

    total_elapsed = time.perf_counter() - t_total
    all_relations = [r for p in papers_data for r in p.get("relations", [])]

    # ── Build graph ────────────────────────────────────────────────────────────
    print(f"\n  Building knowledge graph...")
    graph_out = Path(args.graph_out)
    graph_out.parent.mkdir(parents=True, exist_ok=True)

    try:
        import networkx as nx
        graph = build_graph(papers_data, registry, all_relations)
        print(f"  Nodes : {graph.number_of_nodes()}")
        print(f"  Edges : {graph.number_of_edges()}")

        nx.write_graphml(graph, str(graph_out.with_suffix(".graphml")))
        print(f"  GraphML → {graph_out.with_suffix('.graphml')}")

        graph_json = {
            "nodes": [{"id": n, **d} for n, d in graph.nodes(data=True)],
            "edges": [{"source": u, "target": v, **d} for u, v, d in graph.edges(data=True)],
        }
        graph_out.with_suffix(".json").write_text(
            json.dumps(graph_json, indent=2, ensure_ascii=False)
        )
        print(f"  JSON    → {graph_out.with_suffix('.json')}")

    except ImportError:
        print("  [!] networkx not installed — skipping graph files")
        print("      Install with: pip install networkx")
        graph = None

    # ── Excel ─────────────────────────────────────────────────────────────────
    if args.excel:
        if graph is not None:
            print(f"\n  Generating Excel summary...")
            excel_path = graph_out.parent / "batch_summary.xlsx"
            write_excel(papers_data, registry, graph, excel_path,
                        all_relations=all_relations)
        else:
            print("  [!] Skipping Excel — graph not built (networkx missing)")

    # ── Summary ───────────────────────────────────────────────────────────────
    all_concepts  = registry.concepts()
    cross_paper   = [c for c in all_concepts if c["paper_count"] > 1]

    print(f"\n{'='*60}")
    print(f"  DONE")
    print(f"  Papers processed  : {len(papers_data)}")
    print(f"  Unique concepts   : {len(all_concepts)}")
    print(f"  Cross-paper conc. : {len(cross_paper)}")
    print(f"  Total relations   : {len(all_relations)}")
    if graph is not None:
        print(f"  Graph nodes/edges : {graph.number_of_nodes()} / {graph.number_of_edges()}")
    print(f"  Total time        : {total_elapsed:.1f}s")
    print(f"{'='*60}")

    top = sorted(cross_paper, key=lambda c: (c["paper_count"], c["frequency"]), reverse=True)[:20]
    if top:
        print(f"\n  Top concepts across papers:")
        for i, c in enumerate(top, 1):
            print(f"    {i:2d}. {c['display_text']:<40}  papers={c['paper_count']}  freq={c['frequency']}")


if __name__ == "__main__":
    main()
